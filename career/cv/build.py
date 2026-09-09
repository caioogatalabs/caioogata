#!/usr/bin/env python3
"""
Build compact 1-page DOCX CVs from markdown sources, with REAL hyperlinks.
Uses python-docx for proper hyperlink relationships that work in Word/Pages/Google Docs.
"""
import html
import os
import re
import shutil
import subprocess
import sys
from typing import List

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Resolve the CV directory from this file's location so the build never breaks
# if the repo (or this pillar) moves.
CV_DIR = os.path.dirname(os.path.abspath(__file__))

# Typography
FONT_NAME = "Helvetica Neue"
SIZE_NAME = 16
SIZE_BODY = 9
SIZE_META = 8.5
SIZE_TAGS = 9
SIZE_SECTION = 8
COLOR_TEXT = RGBColor(0x1A, 0x1A, 0x1A)
COLOR_META = RGBColor(0x55, 0x55, 0x55)
COLOR_SECTION = RGBColor(0x44, 0x44, 0x44)
COLOR_LINK_HEX = "0563C1"
COLOR_RULE_HEX = "BBBBBB"


# ---------- Hyperlink helper ----------

def add_hyperlink(paragraph, text: str, url: str, size_pt: float = SIZE_META):
    """Add a real hyperlink (clickable in Word/Pages) to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), FONT_NAME)
    rFonts.set(qn("w:hAnsi"), FONT_NAME)
    rPr.append(rFonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size_pt * 2)))
    rPr.append(sz)

    color = OxmlElement("w:color")
    color.set(qn("w:val"), COLOR_LINK_HEX)
    rPr.append(color)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    run.append(rPr)

    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)

    hyperlink.append(run)
    paragraph._p.append(hyperlink)


# ---------- Markdown line renderer ----------

LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BOLD_RE = re.compile(r"\*\*([^*]+)\*\*")


def render_line(paragraph, text: str, size_pt: float = SIZE_BODY,
                color: RGBColor = COLOR_TEXT, default_bold: bool = False):
    """Render a line that may contain markdown bold and links."""
    tokens = []
    for m in LINK_RE.finditer(text):
        tokens.append(("link", m.start(), m.end(), m.group(1), m.group(2)))
    for m in BOLD_RE.finditer(text):
        if any(t[0] == "link" and t[1] <= m.start() and m.end() <= t[2] for t in tokens):
            continue
        tokens.append(("bold", m.start(), m.end(), m.group(1), None))
    tokens.sort(key=lambda t: t[1])

    pos = 0
    for kind, start, end, content, url in tokens:
        if start > pos:
            chunk = text[pos:start]
            run = paragraph.add_run(chunk)
            _style_run(run, size_pt, color, bold=default_bold)
        if kind == "link":
            add_hyperlink(paragraph, content, url, size_pt=size_pt)
        elif kind == "bold":
            run = paragraph.add_run(content)
            _style_run(run, size_pt, color, bold=True)
        pos = end

    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _style_run(run, size_pt, color, bold=default_bold)


def _style_run(run, size_pt: float, color: RGBColor, bold: bool = False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.bold = bold


# ---------- Markdown parser ----------

def parse_md(md_text: str) -> dict:
    lines = md_text.split("\n")
    out = {
        "name": "",
        "meta_lines": [],
        "tags": "",
        "summary": "",
        "experience": [],
        "education": [],
        "skills": [],
    }

    section = "header"
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if line.startswith("# "):
            out["name"] = line[2:].strip()
            i += 1
            continue

        if line.startswith("## "):
            section = line[3:].strip().lower()
            i += 1
            continue

        if line.strip() == "---" or line.strip() == "":
            i += 1
            continue

        if section == "header":
            stripped = line.strip()
            if (stripped.startswith("**") and stripped.endswith("**")
                    and stripped[2:-2].count("**") == 0):
                out["tags"] = stripped[2:-2]
            else:
                out["meta_lines"].append(stripped)
            i += 1
            continue

        if section == "summary":
            out["summary"] += (" " if out["summary"] else "") + line.strip()
            i += 1
            continue

        if section == "experience":
            header = line.strip()
            desc = ""
            if i + 1 < len(lines) and lines[i+1].strip() and not lines[i+1].strip().startswith("**"):
                desc = lines[i+1].strip()
                i += 2
            else:
                i += 1
            out["experience"].append((header, desc))
            continue

        if section == "education":
            out["education"].append(line.strip())
            i += 1
            continue

        if section == "skills":
            stripped = line.strip()
            if stripped.startswith("**") and stripped.endswith("**"):
                cat = stripped[2:-2]
                if i + 1 < len(lines):
                    items = lines[i+1].strip()
                    out["skills"].append((cat, items))
                    i += 2
                    continue
            i += 1
            continue

        i += 1

    return out


# ---------- DOCX builder ----------

def add_horizontal_rule(paragraph):
    """Add a thin bottom border to a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), COLOR_RULE_HEX)
    pBdr.append(bottom)
    pPr.append(pBdr)


def compress_education(items: List[str]) -> List[str]:
    """Combine multiple Certification lines into one to save space."""
    certs = []
    others = []
    for line in items:
        # Check if this is a certification line
        check = line.lstrip("*").strip().lower()
        if check.startswith("certification"):
            parts = re.split(r"\s+·\s+", line)
            if len(parts) >= 3:
                name = parts[1].strip()
                year = parts[-1].strip()
                certs.append(f"{name} ({year})")
            else:
                certs.append(line)
        else:
            others.append(line)
    if certs:
        others.append(f"**Certifications** · {' · '.join(certs)}")
    return others


def build_docx(parsed: dict, output_path: str):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(SIZE_BODY)
    normal.font.color.rgb = COLOR_TEXT
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15

    # Name
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(parsed["name"])
    run.font.name = FONT_NAME
    run.font.size = Pt(SIZE_NAME)
    run.bold = True
    run.font.color.rgb = COLOR_TEXT

    # Meta lines
    for line in parsed["meta_lines"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        render_line(p, line, size_pt=SIZE_META, color=COLOR_META)

    # Tags
    if parsed["tags"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(parsed["tags"])
        run.font.name = FONT_NAME
        run.font.size = Pt(SIZE_TAGS)
        run.bold = True
        run.font.color.rgb = COLOR_TEXT

    _section_header(doc, "Summary")
    if parsed["summary"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(0)
        render_line(p, parsed["summary"], size_pt=SIZE_BODY)

    _section_header(doc, "Experience")
    for header, desc in parsed["experience"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(0)
        render_line(p, header, size_pt=SIZE_BODY)
        if desc:
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(0)
            render_line(p2, desc, size_pt=SIZE_BODY, color=COLOR_META)

    _section_header(doc, "Education")
    for line in compress_education(parsed["education"]):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(0)
        render_line(p, line, size_pt=SIZE_BODY)

    _section_header(doc, "Skills")
    for cat, items in parsed["skills"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(f"{cat}  ")
        run.font.name = FONT_NAME
        run.font.size = Pt(SIZE_BODY)
        run.bold = True
        render_line(p, items, size_pt=SIZE_BODY, color=COLOR_META)

    doc.save(output_path)


def _section_header(doc, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(title.upper())
    run.font.name = FONT_NAME
    run.font.size = Pt(SIZE_SECTION)
    run.bold = True
    run.font.color.rgb = COLOR_SECTION
    rPr = run._r.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "30")
    rPr.append(spacing)
    add_horizontal_rule(p)


# ---------- HTML builder ----------
# Same markdown source as the DOCX — keeps the .html in lockstep instead of
# being a hand-maintained parallel artifact that silently drifts.

HTML_HEAD = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>__TITLE__</title>
<style>
@page { margin: 0.5in; size: letter; }
* { box-sizing: border-box; }
body {
  font-family: -apple-system, "Helvetica Neue", Helvetica, Arial, sans-serif;
  font-size: 10pt;
  line-height: 1.25;
  color: #1a1a1a;
  margin: 0;
}
h1 {
  font-size: 17pt;
  margin: 0 0 2pt 0;
  font-weight: 700;
  letter-spacing: -0.01em;
}
h2 {
  font-size: 9pt;
  margin: 7pt 0 2pt 0;
  text-transform: uppercase;
  letter-spacing: 0.08em;
  font-weight: 700;
  color: #444;
}
p { margin: 0 0 3pt 0; }
hr {
  border: 0;
  border-top: 0.5pt solid #bbb;
  margin: 5pt 0;
}
.meta { font-size: 9pt; color: #555; margin: 0 0 2pt 0; }
.tags { font-size: 9pt; font-weight: 700; color: #1a1a1a; margin: 3pt 0 0 0; }
.role { font-weight: 700; }
.exp-line { margin: 0 0 4pt 0; }
.skill-cat { font-weight: 700; margin-top: 3pt; margin-bottom: 1pt; }
.skill-items { margin: 0 0 2pt 0; }
a { color: #0a58ca; text-decoration: underline; }
</style>
</head>
<body>"""


def render_html_inline(text: str) -> str:
    """Render markdown bold + links to inline HTML, escaping everything else."""
    tokens = []
    for m in LINK_RE.finditer(text):
        tokens.append(("link", m.start(), m.end(), m.group(1), m.group(2)))
    for m in BOLD_RE.finditer(text):
        if any(t[0] == "link" and t[1] <= m.start() and m.end() <= t[2] for t in tokens):
            continue
        tokens.append(("bold", m.start(), m.end(), m.group(1), None))
    tokens.sort(key=lambda t: t[1])

    out = []
    pos = 0
    for kind, start, end, content, url in tokens:
        if start > pos:
            out.append(html.escape(text[pos:start]))
        if kind == "link":
            out.append(f'<a href="{html.escape(url, quote=True)}">{html.escape(content)}</a>')
        elif kind == "bold":
            out.append(f"<strong>{html.escape(content)}</strong>")
        pos = end
    if pos < len(text):
        out.append(html.escape(text[pos:]))
    return "".join(out)


def build_html(parsed: dict, output_path: str):
    parts = [HTML_HEAD.replace("__TITLE__", html.escape(parsed["name"]))]
    parts.append(f'<h1>{html.escape(parsed["name"])}</h1>')

    for line in parsed["meta_lines"]:
        parts.append(f'<p class="meta">{render_html_inline(line)}</p>')
    if parsed["tags"]:
        parts.append(f'<p class="tags"><strong>{html.escape(parsed["tags"])}</strong></p>')

    parts.append("<hr>")
    parts.append("<h2>Summary</h2>")
    if parsed["summary"]:
        parts.append(f'<p>{render_html_inline(parsed["summary"])}</p>')

    parts.append("<hr>")
    parts.append("<h2>Experience</h2>")
    for header, desc in parsed["experience"]:
        line = render_html_inline(header)
        if desc:
            line += "<br>" + render_html_inline(desc)
        parts.append(f'<p class="exp-line">{line}</p>')

    parts.append("<hr>")
    parts.append("<h2>Education</h2>")
    for line in parsed["education"]:
        parts.append(f'<p class="exp-line">{render_html_inline(line)}</p>')

    parts.append("<hr>")
    parts.append("<h2>Skills</h2>")
    for cat, items in parsed["skills"]:
        parts.append(f'<p class="skill-cat">{render_html_inline(cat)}</p>')
        parts.append(f'<p class="skill-items">{render_html_inline(items)}</p>')

    parts.append("</body>")
    parts.append("</html>")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))


# ---------- PDF export ----------
# Brings the PDF into the pipeline so it regenerates from source instead of
# being a stale, hand-dated artifact. Gracefully skips if LibreOffice is absent.

def export_pdf(docx_path: str, out_dir: str) -> bool:
    """Convert a DOCX to PDF using LibreOffice headless, if available."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
        check=True, capture_output=True,
    )
    return True


# ---------- Driver ----------

def build_cv(md_path: str) -> bool:
    base = os.path.splitext(md_path)[0]
    docx_path = base + ".docx"
    html_path = base + ".html"

    with open(md_path, "r", encoding="utf-8") as f:
        md = f.read()

    parsed = parse_md(md)
    build_docx(parsed, docx_path)
    build_html(parsed, html_path)
    pdf_ok = export_pdf(docx_path, os.path.dirname(os.path.abspath(docx_path)))

    outputs = ".docx + .html" + (" + .pdf" if pdf_ok else " (.pdf skipped: LibreOffice not found)")
    print(f"OK {os.path.basename(base)} -> {outputs}")
    return True


def main():
    md_files = sorted([f for f in os.listdir(CV_DIR) if f.endswith(".md")])
    for f in md_files:
        try:
            build_cv(os.path.join(CV_DIR, f))
        except Exception as e:
            print(f"FAIL {f}: {e}", file=sys.stderr)


if __name__ == "__main__":
    main()
